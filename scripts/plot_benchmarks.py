#!/usr/bin/env python3
"""Plot HolmesGPT XLSX reports from a folder; no server or JSON input needed."""

import argparse
from collections import defaultdict
import csv
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
import hashlib
import importlib.util
import json
import math
import os
from pathlib import Path
import statistics
import sys
from tempfile import TemporaryDirectory
from uuid import uuid4


UTC_PLUS_THREE = timezone(timedelta(hours=3))
NAVY, TEAL, PURPLE, BLUE = "#172B46", "#087F8C", "#7756B5", "#2878B5"


@dataclass
class Run:
    path: Path
    parallel: int
    duration: float
    started: str | None
    model: str
    requests: list[dict]

    def identity(self) -> str:
        if self.started is None:
            # Without a run timestamp, equal measurements need not be the same run.
            return hashlib.sha256(self.path.read_bytes()).hexdigest()
        payload = [self.started, self.parallel, self.duration, self.model, self.requests]
        return hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()


def number(value, label: str, *, optional=False) -> float | None:
    if optional and value in (None, "", "N/A"):
        return None
    try:
        if type(value) not in (int, float) or not math.isfinite(value) or value < 0:
            raise ValueError
        return float(value)
    except (ValueError, OverflowError):
        raise ValueError(f"{label} must be a finite, nonnegative number") from None


def positive_integer(value, label: str) -> int:
    result = number(value, label)
    if result < 1 or not result.is_integer():
        raise ValueError(f"{label} must be a positive integer")
    return int(result)


def read_workbook(path: Path) -> Run:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if not {"Summary", "Results"}.issubset(book.sheetnames):
            raise ValueError("not a HolmesGPT report: Summary and Results sheets are required")
        summary = {str(row[0]).strip(): row[1] for row in book["Summary"].values
                   if len(row) >= 2 and row[0] is not None}
        parallel = positive_integer(summary.get("Parallel requests"), "Parallel requests")
        duration = number(summary.get("Batch duration (s)"), "Batch duration (s)")
        if duration <= 0:
            raise ValueError("Batch duration (s) must be greater than zero")
        started = None
        for label, value in summary.items():
            if label.startswith(("Run started (", "Started (")) and value:
                try:
                    parsed = value if isinstance(value, datetime) else datetime.fromisoformat(str(value))
                    if parsed.tzinfo is not None:
                        started = parsed.astimezone(UTC_PLUS_THREE).isoformat(timespec="milliseconds")
                except ValueError:
                    pass
        # Recover full questions from all historical continuation-row formats.
        questions = defaultdict(str)
        if "Answers" in book.sheetnames:
            for row in book["Answers"].iter_rows(min_row=2, values_only=True):
                if len(row) >= 4 and str(row[1]).startswith("Question") and isinstance(row[3], str):
                    questions[row[0]] += row[3]
        rows = iter(book["Results"].values)
        headers = next(rows, ())
        columns = {str(value).strip(): index for index, value in enumerate(headers) if value is not None}
        required = {"Request", "Outcome", "Completed answer (s)"}
        if not required.issubset(columns):
            raise ValueError("Results is missing required columns: " + ", ".join(sorted(required - columns.keys())))
        requests, ids = [], set()
        for row in rows:
            if all(value is None for value in row):
                continue

            def cell(name):
                index = columns.get(name)
                return row[index] if index is not None and index < len(row) else None

            request_id = positive_integer(cell("Request"), "Request ID")
            if request_id in ids:
                raise ValueError(f"duplicate request ID {request_id}")
            ids.add(request_id)
            outcome = cell("Outcome")
            if not isinstance(outcome, str) or not outcome.strip():
                raise ValueError(f"request {request_id} has no outcome")
            outcome = outcome.strip().lower()
            answer = number(cell("Completed answer (s)"), "Completed answer (s)", optional=outcome != "success")
            first = number(cell("First text (s)"), "First text (s)", optional=True)
            requests.append({"id": request_id, "outcome": outcome,
                             "question": questions.get(request_id, str(cell("Question") or "")),
                             "answer_seconds": answer if outcome == "success" else None,
                             "first_text_seconds": first if outcome == "success" else None})
        if len(requests) != parallel:
            raise ValueError(f"expected {parallel} request rows, found {len(requests)}")
        if summary.get("Total requests") is not None:
            if positive_integer(summary["Total requests"], "Total requests") != len(requests):
                raise ValueError("Total requests does not match Results rows")
        return Run(path, parallel, duration, started, str(summary.get("Requested model") or "Server default"),
                   sorted(requests, key=lambda item: item["id"]))
    finally:
        book.close()


def collect_runs(folder: Path):
    runs, skipped, duplicates, seen = [], [], [], {}
    files = sorted(path for path in folder.rglob("*") if path.is_file() and path.suffix.lower() == ".xlsx"
                   and not path.name.startswith("~$") and not path.name.lower().endswith(".tmp.xlsx"))
    for path in files:
        try:
            run = read_workbook(path)
            identity = run.identity()
            if identity in seen:
                duplicates.append({"file": str(path), "same_run_as": str(seen[identity])})
                continue
            seen[identity] = path
            runs.append(run)
        except Exception as exc:
            skipped.append({"file": str(path), "reason": str(exc)})
    return runs, skipped, duplicates


def aggregate(runs: list[Run]) -> list[dict]:
    groups = defaultdict(list)
    for run in runs:
        groups[run.parallel].append(run)
    metrics = []
    for parallel, group in sorted(groups.items()):
        requests = [result for run in group for result in run.requests]
        successes = [result for result in requests if result["outcome"] == "success"]
        answers = sorted(result["answer_seconds"] for result in successes)
        first = [result["first_text_seconds"] for result in successes if result["first_text_seconds"] is not None]
        duration = math.fsum(run.duration for run in group)
        metrics.append({
            "parallel": parallel, "runs": len(group), "requests": len(requests), "successes": len(successes),
            "failures": len(requests) - len(successes),
            "mean_answer_seconds": statistics.fmean(answers) if answers else None,
            "p95_answer_seconds": answers[math.ceil(0.95 * len(answers)) - 1] if answers else None,
            "mean_first_text_seconds": statistics.fmean(first) if first else None,
            "first_text_samples": len(first), "batch_seconds": duration,
            "successful_requests_per_second": len(successes) / duration,
            "success_rate_percent": len(successes) / len(requests) * 100,
        })
    return metrics


PANELS = {
    "answer_latency": ("Answer time vs parallelism", "Seconds · lower is better",
                       [("mean_answer_seconds", "Average", TEAL, "-"),
                        ("p95_answer_seconds", "p95", PURPLE, "--")]),
    "first_text_latency": ("Time until the first assistant text", "Seconds · lower is better",
                           [("mean_first_text_seconds", "Average first text", BLUE, "-")]),
    "throughput": ("Successful answers per second", "Answers / second · higher is better",
                   [("successful_requests_per_second", "Throughput", TEAL, "-")]),
    "success_rate": ("Success rate vs parallelism", "Percent · higher is better",
                     [("success_rate_percent", "Success rate", BLUE, "-")]),
}


def draw_panel(ax, name: str, metrics: list[dict], *, compact=False):
    title, ylabel, series = PANELS[name]
    xs = [row["parallel"] for row in metrics]
    any_values = False
    for key, label, color, linestyle in series:
        ys = [row[key] if row[key] is not None else math.nan for row in metrics]
        if any(math.isfinite(value) for value in ys):
            any_values = True
            ax.plot(xs, ys, label=label, color=color, linestyle=linestyle, linewidth=2.6,
                    marker="o", markersize=6, markeredgecolor="white", markeredgewidth=1.2, zorder=3)
    ax.set_title(title, loc="left", fontsize=15 if compact else 20, fontweight="bold", pad=22, color=NAVY)
    ax.set_xlabel("Parallel requests", labelpad=12)
    ax.set_ylabel(ylabel, labelpad=12)
    ax.set_axisbelow(True)
    ax.grid(axis="y", color="#DCE5EE", linewidth=0.8)
    ax.spines[["top", "right"]].set_visible(False)
    ax.spines[["left", "bottom"]].set_color("#CBD5E1")
    ax.tick_params(colors="#526277", length=0, pad=8)
    tick_step = max(1, math.ceil(len(xs) / 10))
    ax.set_xticks(sorted(set(xs[::tick_step] + [xs[-1]])))
    padding = max(0.5, (xs[-1] - xs[0]) * 0.07)
    ax.set_xlim(max(0, xs[0] - padding), xs[-1] + padding)
    ax.set_ylim(bottom=0)
    if name == "success_rate":
        ax.set_ylim(0, 114)
        ax.set_yticks([0, 25, 50, 75, 100], ["0%", "25%", "50%", "75%", "100%"])
    else:
        ax.set_ylim(top=max(0.001, ax.get_ylim()[1] * 1.25))
    if any_values:
        if len(series) > 1:
            ax.legend(loc="upper left", frameon=False, ncol=2)
        if len(xs) <= (6 if compact else 10):
            for row in metrics:
                y = row[series[0][0]]
                if y is None:
                    continue
                label_y = y
                if name == "answer_latency":
                    label = f"Avg {y:.3g}s\nn={row['successes']}"
                    label_y = max(y, row["p95_answer_seconds"])
                elif name == "first_text_latency":
                    label = f"{y:.3g}s\nn={row['first_text_samples']}"
                elif name == "success_rate":
                    label = f"{y:.3g}%\n{row['successes']}/{row['requests']}"
                else:
                    label = f"{y:.3g}"
                ax.annotate(label, (row["parallel"], label_y), xytext=(0, 13), textcoords="offset points",
                            ha="center", va="bottom", fontsize=9, color=NAVY,
                            bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.85, "pad": 1.5})
    else:
        ax.text(0.5, 0.5, "No successful timing samples available", transform=ax.transAxes,
                ha="center", va="center", color="#64748B", fontsize=12)


def export_graphs(metrics: list[dict], directory: Path, generated_at: str):
    # Keep standalone use headless and usable on hosts with read-only home folders.
    with TemporaryDirectory(prefix="holmes-plot-cache-") as cache:
        previous = os.environ.get("MPLCONFIGDIR")
        if previous is None:
            os.environ["MPLCONFIGDIR"] = cache
        try:
            _render_graphs(metrics, directory, generated_at)
        finally:
            if previous is None:
                os.environ.pop("MPLCONFIGDIR", None)


def _render_graphs(metrics: list[dict], directory: Path, generated_at: str):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    runs = sum(row["runs"] for row in metrics)
    requests = sum(row["requests"] for row in metrics)
    levels = "parallelism level" if len(metrics) == 1 else "parallelism levels"
    subtitle = f"{runs} unique runs  /  {requests} requests  /  {len(metrics)} {levels}"
    footnote = ("Timing: successful requests only; n = available samples. p95: nearest rank, not an uncertainty band.\n"
                "Throughput = successful answers / summed batch durations. All workbooks are pooled. Generated " + generated_at)
    style = {"font.family": "DejaVu Sans", "font.size": 11, "text.color": NAVY,
             "axes.labelcolor": "#526277", "figure.facecolor": "#F5F8FC", "axes.facecolor": "white",
             "savefig.facecolor": "#F5F8FC", "svg.fonttype": "none", "text.usetex": False}
    with plt.rc_context(style):
        for name in [*PANELS, "overview"]:
            if name == "overview":
                fig, axes = plt.subplots(2, 2, figsize=(16, 11))
                fig.subplots_adjust(left=0.08, right=0.96, top=0.82, bottom=0.14, hspace=0.58, wspace=0.30)
                for ax, panel in zip(axes.flat, PANELS):
                    draw_panel(ax, panel, metrics, compact=True)
            else:
                fig, ax = plt.subplots(figsize=(12, 7.5))
                fig.subplots_adjust(left=0.11, right=0.94, top=0.74, bottom=0.19)
                draw_panel(ax, name, metrics)
            try:
                fig.text(0.08, 0.94, "HOLMESGPT  /  BENCHMARK ANALYSIS", fontsize=20, fontweight="bold", color=NAVY)
                fig.text(0.08, 0.90, subtitle, fontsize=12, color="#526277")
                fig.text(0.08, 0.035, footnote, fontsize=8.5, color="#526277", linespacing=1.8)
                for extension in ("png", "svg"):
                    fig.savefig(directory / f"{name}.{extension}", dpi=180,
                                metadata={"Title": "HolmesGPT benchmark analysis", "Date": generated_at})
            finally:
                plt.close(fig)


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__,
        epilog="Recursively reads .xlsx files. Repeated runs are pooled by parallelism; copies of a run are counted once.")
    parser.add_argument("folder", type=Path, help="folder containing benchmark XLSX files (subfolders included)")
    parser.add_argument("--output-dir", "-o", type=Path, help="parent folder for graphs (default: INPUT/graphs)")
    args = parser.parse_args(argv)
    folder = args.folder.expanduser().resolve()
    if not folder.is_dir():
        parser.error("folder must be an existing directory")
    if any(importlib.util.find_spec(name) is None for name in ("openpyxl", "matplotlib")):
        print('Missing chart dependencies. Run: python -m pip install -e ".[charts]"', file=sys.stderr)
        return 1
    print(f"Reading benchmark workbooks from {folder}…", flush=True)
    runs, skipped, duplicates = collect_runs(folder)
    for item in skipped:
        print("Skipped " + " ".join(item["file"].split()) + ": " + " ".join(item["reason"].split()), file=sys.stderr)
    if not runs:
        print("No valid benchmark workbooks found. Expected Summary and Results sheets.", file=sys.stderr)
        return 1
    notices = []
    if len({run.model for run in runs}) > 1:
        notices.append("Inputs contain different requested models; all files are pooled.")
    if len({tuple(sorted({r['question'] for r in run.requests})) for run in runs}) > 1:
        notices.append("Inputs contain different question sets; all files are pooled.")
    if len({run.parallel for run in runs}) == 1:
        notices.append("Only one parallelism level is present; add runs with other --parallel values to see a scaling curve.")
    for notice in notices:
        print("Note: " + notice, file=sys.stderr)
    generated = datetime.now(UTC_PLUS_THREE)
    parent = args.output_dir.expanduser().resolve() if args.output_dir else folder / "graphs"
    output = parent / (generated.strftime("%Y%m%dT%H%M%S%f%z") + "-" + uuid4().hex[:8])
    try:
        metrics = aggregate(runs)
        output.mkdir(parents=True)
        with (output / "metrics.csv").open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(metrics[0]))
            writer.writeheader()
            writer.writerows(metrics)
        manifest = {
            "generated_at": generated.isoformat(timespec="seconds"), "timezone": "UTC+3",
            "included": [{"file": str(run.path), "started_at": run.started, "parallel": run.parallel,
                          "requests": len(run.requests), "requested_model": run.model} for run in runs],
            "skipped": skipped, "duplicates": duplicates, "notices": notices,
            "methods": {"latency": "Pooled successful request measurements, not averages of run averages.",
                        "p95": "Nearest rank among successful requests at each parallelism.",
                        "throughput": "Sum of successful requests / sum of batch durations.",
                        "missing": "Unavailable measurements are blank in CSV and omitted from plots."},
        }
        (output / "sources.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")
        export_graphs(metrics, output, manifest["generated_at"])
    except Exception as exc:
        print(f"Graph export failed: {exc}. Any partial outputs are in {output}", file=sys.stderr)
        return 1
    print(f"Loaded {len(runs)} runs; ignored {len(duplicates)} duplicate exports; skipped {len(skipped)} invalid files.")
    print(f"Graphs: {output}\nStart with overview.png. PNG, SVG, metrics.csv, and sources.json saved.")
    return 1 if skipped else 0


if __name__ == "__main__":
    raise SystemExit(main())
