"""Simple, filterable Excel sheets with literal response text."""

import math
from pathlib import Path
import re


INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")


def _chunks(value: str):
    text = INVALID_XML.sub("", value)
    start, units = 0, 0
    for index, char in enumerate(text):
        size = 2 if ord(char) > 0xFFFF else 1
        if units + size > 32000:
            yield text[start:index]
            start, units = index, 0
        units += size
    yield text[start:]


def write_excel(report: dict, directory: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    results = book.create_sheet("Results")
    answers = book.create_sheet("Answers")
    header_fill = PatternFill("solid", fgColor="24445C")
    stripe_fill = PatternFill("solid", fgColor="F3F5F7")

    def append(sheet, values):
        row = sheet.max_row + 1 if sheet.cell(1, 1).value is not None else 1
        for column, value in enumerate(values, 1):
            value = "N/A" if value is None else value
            if isinstance(value, str):
                clean = INVALID_XML.sub("", value)
                value = next(_chunks(clean))
                if len(value) < len(clean):
                    value = value[:1500] + "… [full text in Answers / JSON]"
            cell = sheet.cell(row, column, value)
            if isinstance(value, str):
                cell.data_type = "s"  # Never interpret a response as a formula.
            cell.font = Font(name="Calibri", size=11, color="243746")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(value, float):
                cell.number_format = "0.000"
            if row > 1 and row % 2 == 0:
                cell.fill = stripe_fill
        return row

    def format_sheet(sheet, widths, filters=False):
        sheet.freeze_panes = "A2"
        sheet.row_dimensions[1].height = 30
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = header_fill
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        if filters:
            # A worksheet filter is sufficient; do not also define an Excel table.
            sheet.auto_filter.ref = sheet.dimensions

    stats = report["summary"]
    append(summary, ["Metric", "Value"])
    metrics = [
        ("Total requests", stats["total_requests"]),
        ("Successful requests", stats["successful_requests"]),
        ("Unsuccessful requests", stats["unsuccessful_requests"]),
        ("Success rate", stats["success_rate"]),
        ("Batch duration (s)", stats["batch_duration_seconds"]),
        ("Successful requests / second", stats["successful_requests_per_second"]),
        ("Completed answer sample count", stats["completion_seconds"]["count"]),
        *((f"{key.title()} completed answer (s)", stats["completion_seconds"][key])
          for key in ("min", "mean", "median", "p95", "max")),
        ("Endpoint", report["settings"]["api_url"]),
        ("Requested model", report["settings"]["requested_model"] or "Server default"),
        ("Parallel requests", report["settings"]["parallel"]),
        ("Request deadline (s)", report["settings"]["timeout_seconds"]),
        ("Started (UTC)", report["started_at"]),
        ("Finished (UTC)", report["finished_at"]),
        ("Interrupted", report["interrupted"]),
    ]
    for label, value in metrics:
        row = append(summary, [label, value])
        if label == "Success rate":
            summary.cell(row, 2).number_format = "0.0%"
    for name, note in report["measurement_notes"].items():
        row = append(summary, [name.replace("_", " ").title() + " note", note])
        summary.row_dimensions[row].height = 42
    format_sheet(summary, [36, 90])

    append(results, ["Request", "Question", "Outcome", "First text (s)", "Completed answer (s)",
                     "Elapsed (s)", "Tool calls", "Tool errors", "Last prompt tokens",
                     "Last output tokens", "Last total tokens", "Requested model", "HTTP status", "Error"])
    append(answers, ["Request", "Content", "Part", "Text"])
    for result in report["requests"]:
        usage = result["last_reported_usage"] or {}

        def token(name):
            value = usage.get(name)
            try:
                return value if type(value) in (int, float) and math.isfinite(value) and value >= 0 else None
            except OverflowError:
                return None

        row = append(results, [
            result["request_id"], result["question"], result["outcome"],
            result["first_text_seconds"], result["completion_seconds"], result["elapsed_seconds"],
            result["tool_call_count"], result["tool_error_count"], token("prompt_tokens"),
            token("completion_tokens"), token("total_tokens"), result["requested_model"] or "Server default",
            result["http_status"], result["error"],
        ])
        results.cell(row, 3).font = Font(name="Calibri", size=11, bold=True,
                                        color="24713C" if result["outcome"] == "success" else "A02C2C")
        results.row_dimensions[row].height = 42
        for kind, text in (("Question", result["question"]), ("Answer", result["answer"]), ("Error", result["error"])):
            if text is None:
                continue
            for part, chunk in enumerate(_chunks(text), 1):
                row = append(answers, [result["request_id"], kind, part, chunk])
                lines = sum(max(1, math.ceil(len(line) / 95)) for line in chunk.split("\n"))
                answers.row_dimensions[row].height = min(300, max(30, lines * 15))
    format_sheet(results, [10, 48, 22, 18, 22, 18, 14, 14, 19, 19, 19, 24, 14, 48], filters=True)
    format_sheet(answers, [10, 14, 10, 100], filters=True)

    destination = directory / "report.xlsx"
    temporary = directory / "report.tmp.xlsx"
    try:
        book.save(temporary)
        temporary.replace(destination)
    finally:
        book.close()
    return destination
