"""Graphs use the actual workbook contract, not cached summary averages."""

import csv
import json
from pathlib import Path
import subprocess
import sys

from openpyxl import Workbook
import pytest


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "plot_benchmarks.py"


def workbook(path, times, *, duration=10, started="2026-10-18T12:00:00+03:00", legacy=False):
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Summary"
    sheet.append(["HolmesGPT", "Report"])
    for row in [("Parallel requests", len(times)), ("Total requests", len(times)),
                ("Batch duration (s)", duration), ("Requested model", "Server default"),
                ("Started (UTC)" if legacy else "Run started (UTC+3)", started),
                ("Mean completed answer (s)", 999)]:
        sheet.append(row)
    results = book.create_sheet("Results")
    headers = ["Request", "Question", "Outcome", "First text (s)", "Completed answer (s)", "Last output tokens"]
    if legacy:
        headers.reverse()  # Header names must determine the schema, not positions.
    results.append(headers)
    for i, value in enumerate(times, 1):
        record = dict(zip(["Request", "Question", "Outcome", "First text (s)", "Completed answer (s)", "Last output tokens"],
                          [i, "Check pods", "success" if value is not None else "timeout",
                           value / 10 if value is not None else "N/A", value if value is not None else "N/A", 100]))
        results.append([record[key] for key in headers])
    book.save(path)
    book.close()
    return path


def test_pooling_uses_requests_excludes_failures_and_weights_throughput(tmp_path):
    from scripts.plot_benchmarks import read_workbook, aggregate

    first = read_workbook(workbook(tmp_path / "one.xlsx", [2, 4, None], duration=10))
    second = read_workbook(workbook(tmp_path / "two.xlsx", [8, None, None], duration=20, legacy=True))
    row = aggregate([first, second])[0]
    assert row["parallel"] == 3
    assert row["runs"] == 2 and row["requests"] == 6 and row["successes"] == 3
    assert row["mean_answer_seconds"] == pytest.approx(14 / 3)
    assert row["p95_answer_seconds"] == 8
    assert row["mean_first_text_seconds"] == pytest.approx(1.4 / 3)
    assert row["first_text_samples"] == 3
    assert row["successful_requests_per_second"] == pytest.approx(3 / 30)
    assert row["success_rate_percent"] == 50


def test_recursive_scan_deduplicates_same_run_across_layouts_and_offsets(tmp_path):
    from scripts.plot_benchmarks import collect_runs

    workbook(tmp_path / "original" / "report.xlsx", [1, 3])
    workbook(tmp_path / "copy.xlsx", [1, 3], started="2026-10-18T09:00:00Z", legacy=True)
    workbook(tmp_path / "new-run.xlsx", [1, 3], started="2026-10-18T13:00:00+03:00")
    (tmp_path / "~$report.xlsx").write_bytes(b"Excel lock file")
    (tmp_path / "broken.xlsx").write_bytes(b"broken zip")
    runs, skipped, duplicates = collect_runs(tmp_path)
    assert len(runs) == 2
    assert len(skipped) == 1 and "broken.xlsx" in skipped[0]["file"]
    assert len(duplicates) == 1


def test_no_successes_remain_unavailable_not_zero(tmp_path):
    from scripts.plot_benchmarks import read_workbook, aggregate

    row = aggregate([read_workbook(workbook(tmp_path / "failed.xlsx", [None, None]))])[0]
    assert row["mean_answer_seconds"] is None
    assert row["p95_answer_seconds"] is None
    assert row["mean_first_text_seconds"] is None
    assert row["success_rate_percent"] == 0
    assert row["successful_requests_per_second"] == 0


def test_p95_nearest_rank_and_missing_first_text_samples(tmp_path):
    from openpyxl import load_workbook
    from scripts.plot_benchmarks import read_workbook, aggregate

    path = workbook(tmp_path / "twenty.xlsx", list(range(1, 21)), duration=25)
    book = load_workbook(path)
    book["Results"]["D2"] = "N/A"
    book.save(path)
    book.close()
    row = aggregate([read_workbook(path)])[0]
    assert row["p95_answer_seconds"] == 19
    assert row["mean_answer_seconds"] == 10.5
    assert row["first_text_samples"] == 19
    assert row["mean_first_text_seconds"] == pytest.approx(1.1)


@pytest.mark.parametrize("cell,value", [("E2", None), ("E2", -1), ("A2", 0)])
def test_invalid_request_measurements_rejected(tmp_path, cell, value):
    from openpyxl import load_workbook
    from scripts.plot_benchmarks import read_workbook

    path = workbook(tmp_path / "bad.xlsx", [1])
    book = load_workbook(path)
    book["Results"][cell] = value
    book.save(path)
    book.close()
    with pytest.raises(ValueError):
        read_workbook(path)


def test_standalone_script_exports_graphs_and_reproducible_values(tmp_path):
    workbook(tmp_path / "inputs" / "one.xlsx", [2])
    workbook(tmp_path / "inputs" / "two.xlsx", [3, 5], started="2026-10-18T13:00:00+03:00")
    result = subprocess.run([sys.executable, str(SCRIPT), str(tmp_path / "inputs"),
                             "--output-dir", str(tmp_path / "charts")],
                            cwd=tmp_path, capture_output=True, text=True, timeout=60)
    assert result.returncode == 0, result.stderr
    directories = list((tmp_path / "charts").iterdir())
    assert len(directories) == 1
    output = directories[0]
    for name in ("answer_latency", "first_text_latency", "throughput", "success_rate", "overview"):
        assert (output / f"{name}.png").read_bytes().startswith(b"\x89PNG\r\n\x1a\n")
        assert "<svg" in (output / f"{name}.svg").read_text()
    with (output / "metrics.csv").open() as handle:
        rows = list(csv.DictReader(handle))
    assert [int(row["parallel"]) for row in rows] == [1, 2]
    assert [float(row["mean_answer_seconds"]) for row in rows] == [2, 4]
    manifest = json.loads((output / "sources.json").read_text())
    assert manifest["generated_at"].endswith("+03:00")
    assert len(manifest["included"]) == 2
    assert manifest["skipped"] == []


def test_cli_keeps_all_failed_runs_and_reports_invalid_files(tmp_path):
    workbook(tmp_path / "inputs" / "failed.xlsx", [None, None])
    (tmp_path / "inputs" / "invalid.xlsx").write_bytes(b"broken")
    result = subprocess.run([sys.executable, str(SCRIPT), str(tmp_path / "inputs")],
                            capture_output=True, text=True, timeout=60)
    assert result.returncode == 1
    assert "invalid.xlsx" in result.stderr
    output = next((tmp_path / "inputs" / "graphs").iterdir())
    with (output / "metrics.csv").open() as handle:
        row = next(csv.DictReader(handle))
    assert row["mean_answer_seconds"] == ""
    assert float(row["success_rate_percent"]) == 0
    assert (output / "overview.png").exists()


def test_empty_input_and_invalid_folder_do_not_create_outputs(tmp_path):
    for directory, code in ((tmp_path, 1), (tmp_path / "missing", 2)):
        result = subprocess.run([sys.executable, str(SCRIPT), str(directory)],
                                capture_output=True, text=True, timeout=20)
        assert result.returncode == code
    assert not (tmp_path / "graphs").exists()
