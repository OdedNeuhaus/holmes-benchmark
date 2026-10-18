import json
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from holmes_benchmark.config import Config
from holmes_benchmark.runner import Batch, RequestResult


def sample_batch():
    return Batch([
        RequestResult(1, "A", outcome="success", error=None, completion_seconds=1,
                      first_text_seconds=0.1, elapsed_seconds=1.01, answer="Answer A"),
        RequestResult(2, "B", outcome="success", error=None, completion_seconds=3,
                      first_text_seconds=0.2, elapsed_seconds=3.01, answer="Answer B"),
        RequestResult(3, "A", outcome="timeout", error="Timed out", elapsed_seconds=10),
    ], "2026-10-18T07:00:00+00:00", "2026-10-18T07:00:10+00:00", 10)


def test_summary_uses_only_successes_and_nearest_rank_p95():
    from holmes_benchmark.report import build_report

    report = build_report(Config("http://host/api/chat", ("A", "B"), parallel=3), sample_batch())
    summary = report["summary"]
    assert summary["total_requests"] == 3
    assert summary["successful_requests"] == 2
    assert summary["success_rate"] == pytest.approx(2 / 3)
    assert summary["successful_requests_per_second"] == 0.2
    assert summary["completion_seconds"] == {
        "count": 2, "min": 1, "mean": 2, "median": 2, "p95": 3, "max": 3,
    }
    assert report["schema_version"] == 1
    assert report["requests"][2]["completion_seconds"] is None


def test_no_success_is_na_not_zero(capsys):
    from holmes_benchmark.report import build_report, print_summary

    batch = sample_batch()
    batch.requests = batch.requests[2:]
    report = build_report(Config("http://host/api/chat", ("A",)), batch)
    assert report["summary"]["completion_seconds"] == {
        "count": 0, "min": None, "mean": None, "median": None, "p95": None, "max": None,
    }
    print_summary(report)
    output = capsys.readouterr().out
    assert "timeout" in output and "N/A" in output
    assert "Timed out" in output


def test_terminal_summary_shows_usage_snapshot_and_handles_invalid_metrics():
    from io import StringIO
    from rich.console import Console
    from holmes_benchmark.report import build_report, print_summary

    batch = sample_batch()
    batch.requests[0].question = "Check [bold]pods\x1b[2J secret"
    batch.requests[0].last_reported_usage = {"prompt_tokens": 0, "completion_tokens": 9, "total_tokens": 9}
    batch.requests[1].last_reported_usage = {"prompt_tokens": True, "completion_tokens": -2, "total_tokens": "secret"}
    report = build_report(Config("http://host/api/chat", ("A",), token="secret"), batch)
    stream = StringIO()
    print_summary(report, console=Console(file=stream, width=100, force_terminal=False))
    output = stream.getvalue()
    assert "0 / 9 / 9" in output
    assert "N/A / N/A / N/A" in output
    assert "timeout" in output and "10.000" in output
    assert "not investigation totals" in output
    assert "secret" not in output and "\x1b" not in output
    assert "[bold]pods" in output and "[REDACTED]" in output
    assert "Success: 2/3" in output and "Completed answers (n=2)" in output


def test_json_preserves_details_and_redacts_credential_everywhere(tmp_path):
    from holmes_benchmark.report import build_report, write_json

    secret = 's3cr"et\\key'
    batch = sample_batch()
    batch.requests[0].answer = f"Echoed {secret}\x00 שלום"
    batch.requests[0].raw_events = [{
        "data": {"nested": [secret]}, "raw_data": json.dumps({"echo": secret}),
        "raw_lines": [f"data: {json.dumps({'echo': secret})}"],
    }]
    batch.requests[1].error = f"Error: {secret}"
    report = build_report(Config("http://host/api/chat", ("A",), token=secret), batch)
    write_json(report, tmp_path)
    saved = json.loads((tmp_path / "results.json").read_text())
    assert "token" not in saved["settings"]
    assert saved["requests"][0]["answer"] == "Echoed [REDACTED]\x00 שלום"
    assert saved["requests"][0]["raw_events"][0]["data"]["nested"] == ["[REDACTED]"]
    assert "s3cr" not in (tmp_path / "results.json").read_text()
    assert secret in batch.requests[0].answer  # Export sanitization does not mutate measurements.


def test_excel_simple_filters_literal_text_and_long_answer_chunks(tmp_path):
    from holmes_benchmark.report import build_report, write_excel

    batch = sample_batch()
    batch.requests[0].question = '=HYPERLINK("https://example.test","click")'
    batch.requests[0].answer = "=1+1\x00\uffff" + "שלום😀" * 12000
    batch.requests[0].last_reported_usage = {"prompt_tokens": 12, "completion_tokens": 4, "total_tokens": 16}
    report = build_report(Config("http://host/api/chat", ("A",)), batch)
    write_excel(report, tmp_path)
    path = tmp_path / "report.xlsx"
    book = load_workbook(path)
    assert book.sheetnames == ["Summary", "Results", "Answers"]
    summary, results, answers = (book[name] for name in book.sheetnames)
    assert not summary._charts
    assert all(not sheet.merged_cells.ranges for sheet in book)
    assert results.freeze_panes == "A2"
    assert not results.tables and not answers.tables
    assert results.auto_filter.ref == results.dimensions
    assert answers.auto_filter.ref == answers.dimensions
    headers = {cell.value: cell.column for cell in results[1]}
    q = results.cell(2, headers["Question"])
    assert q.value.startswith("=HYPERLINK") and q.data_type == "s"
    assert results.cell(2, headers["Last total tokens"]).value == 16
    assert results.cell(3, headers["Last total tokens"]).value == "N/A"
    assert not results.conditional_formatting
    chunks = [row[3].value for row in answers.iter_rows(min_row=2)
              if row[0].value == 1 and row[1].value.startswith("Answer")]
    assert len(chunks) > 1
    assert "".join(chunks) == "=1+1" + "שלום😀" * 12000
    assert all(len(s.encode("utf-16-le")) // 2 <= 32767 for s in chunks)
    assert all(cell.data_type != "f" for row in answers for cell in row)
    book.close()
    with ZipFile(path) as archive:
        assert b"<f>" not in archive.read("xl/worksheets/sheet2.xml")
        # Keep one standard worksheet filter; no table/drawing relationships
        # for Excel to repair or graphics to distort the layout.
        assert not any(name.startswith(("xl/tables/", "xl/charts/", "xl/drawings/"))
                       for name in archive.namelist())
        import xml.etree.ElementTree as ET
        ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        for name in ("xl/worksheets/sheet2.xml", "xl/worksheets/sheet3.xml"):
            root = ET.fromstring(archive.read(name))
            assert len(root.findall("m:autoFilter", ns)) == 1
            assert root.find("m:tableParts", ns) is None


def test_excel_all_failed_still_writes_usable_report(tmp_path):
    from holmes_benchmark.report import build_report, write_excel

    batch = sample_batch()
    batch.requests = batch.requests[2:]
    write_excel(build_report(Config("http://host/api/chat", ("A",)), batch), tmp_path)
    book = load_workbook(tmp_path / "report.xlsx")
    assert book["Results"].max_row == 2
    assert "N/A" in [cell.value for row in book["Summary"] for cell in row]
    book.close()


def test_excel_keeps_one_answer_in_one_cell_and_preserves_newlines(tmp_path):
    from holmes_benchmark.report import build_report, write_excel

    batch = sample_batch()
    answer = "Overview\n" + "A long explanation of the pod status. " * 200 + "\n" * 40 + "Final note."
    batch.requests[0].answer = answer
    batch.requests[0].question = "Why are these pods restarting? " * 40
    write_excel(build_report(Config("http://host/api/chat", ("A",)), batch), tmp_path)
    book = load_workbook(tmp_path / "report.xlsx")
    sheet = book["Answers"]
    records = [row for row in sheet.iter_rows(min_row=2)
               if row[0].value == 1 and row[1].value == "Answer"]
    assert len(records) == 1
    assert "".join(row[3].value or "" for row in records) == answer
    assert records[0][2].value == "Complete"
    for row in records:
        cell = row[3]
        assert cell.alignment.wrap_text
        assert 30 <= sheet.row_dimensions[cell.row].height <= 409
    assert len(book["Results"]["B2"].value) < 250
    assert "Answers" in book["Results"]["B2"].value
    book.close()


@pytest.mark.parametrize("answer, count", [("a" * 32767, 1), ("a" * 32768, 2), ("line\n" * 254, 2)],
                         ids=["at-character-limit", "over-character-limit", "over-newline-limit"])
def test_excel_splits_only_at_storage_limits_and_labels_continuations(tmp_path, answer, count):
    from holmes_benchmark.report import build_report, write_excel

    batch = sample_batch()
    batch.requests = batch.requests[:1]
    batch.requests[0].answer = answer
    write_excel(build_report(Config("http://host/api/chat", ("A",)), batch), tmp_path)
    book = load_workbook(tmp_path / "report.xlsx")
    rows = [row for row in book["Answers"].iter_rows(min_row=2)
            if row[1].value.startswith("Answer")]
    assert len(rows) == count
    assert "".join(row[3].value for row in rows) == answer
    assert all(len(row[3].value.encode("utf-16-le")) // 2 <= 32767 for row in rows)
    assert all(row[3].value.count("\n") <= 253 for row in rows)
    if count > 1:
        assert rows[1][1].value == "Answer (continued)"
        assert [row[2].value for row in rows] == ["1/2", "2/2"]
    book.close()


@pytest.mark.parametrize("secret", ["model", "success", "text", "a"])
def test_redaction_preserves_schema_and_outcomes_but_cleans_provider_keys(tmp_path, capsys, secret):
    from holmes_benchmark.report import build_report, print_summary, write_excel, write_json

    batch = sample_batch()
    batch.requests[0].answer = f"Echo {secret}"
    batch.requests[0].metadata = {secret: {"echo": secret}}
    batch.requests[0].raw_events = [{
        "event": "ai_answer_end", "id": None, "complete": True,
        "received_at": "2026-10-18T07:00:00+00:00", "elapsed_seconds": 1,
        "data": {secret: secret}, "raw_data": json.dumps({secret: secret}),
        "raw_lines": [f"data: {json.dumps({secret: secret})}"],
    }]
    report = build_report(Config("http://host/api/chat", ("A",), token=secret), batch)
    assert report["summary"]["successful_requests"] == 2
    assert report["requests"][0]["outcome"] == "success"
    assert report["requests"][0]["first_text_seconds"] == 0.1
    assert report["settings"]["requested_model"] is None
    assert report["requests"][0]["answer"] == "Echo [REDACTED]"
    assert report["requests"][0]["metadata"] == {"[REDACTED]": {"echo": "[REDACTED]"}}
    assert report["requests"][0]["raw_events"][0]["data"] == {"[REDACTED]": "[REDACTED]"}
    write_json(report, tmp_path)
    print_summary(report)
    write_excel(report, tmp_path)
    assert "Success: 2/3" in capsys.readouterr().out
    assert (tmp_path / "report.xlsx").exists()
